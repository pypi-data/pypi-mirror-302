from __future__ import annotations

from dataclasses import dataclass

from docx import Document as init_docx
from docx.document import Document as DocxDocument
from fitz import Document as PDFDocument
from fitz import open as init_pdf
from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook as init_xlsx
from pptx import Presentation as init_pptx
from pptx.presentation import Presentation as PPTXDocument

from ._proxy import DocumentFile


@dataclass
class DocxFile(DocumentFile[DocxDocument]):
    def __load__(self):
        return init_docx(self.name)

    def extract_text(self):
        doc = self.__load__()
        for paragraph in doc.paragraphs:
            if paragraph.text:
                yield paragraph.text
            else:
                continue

    def extract_image(self):
        doc = self.__load__()
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    continue
                else:
                    for inline in run.element.iter():  # type: ignore
                        if inline.tag.endswith("inline"):  # type: ignore
                            for pic in inline.iter():  # type: ignore
                                if pic.tag.endswith("blip"):  # type: ignore
                                    image = pic.embed  # type: ignore
                                    image_part = run.part.related_parts[image]
                                    yield image_part.blob
                                else:
                                    continue
                        else:
                            continue


@dataclass
class PDFFile(DocumentFile[PDFDocument]):
    def __load__(self):
        return init_pdf(self.name)

    def extract_text(self):  # type: ignore
        text_doc = self.__load__()
        for page in text_doc:
            yield page.get_textpage().extractTEXT(sort=True)

    def extract_image(self):
        img_doc = self.__load__()
        for page in img_doc:  # type: ignore
            for img in page.get_images():  # type: ignore
                xref = img[0]  # type: ignore
                base_image = img_doc.extract_image(xref)  # type: ignore
                image_bytes = base_image["image"]  # type: ignore
                assert isinstance(image_bytes, bytes)
                yield image_bytes


@dataclass
class ExcelFile(DocumentFile[Workbook]):
    def __load__(self):
        return init_xlsx(self.name)


@dataclass
class PPTXFile(DocumentFile[PPTXDocument]):
    def __load__(self):
        return init_pptx(self.name)
